"""
Kaivor Knowledge Indexer
"""

from pathlib import Path

from docx import Document
from pypdf import PdfReader

from core.knowledge.document_index import DocumentIndex
from core.knowledge.indexed_document import IndexedDocument


class KnowledgeIndexer:
    """Indexes supported knowledge documents."""

    def __init__(self):
        self.index = DocumentIndex()

    def build(self, directory="knowledge"):
        """Build an in-memory searchable document index."""

        self.index.clear()

        root = Path(directory)

        if not root.exists():
            return self.index

        for file in root.rglob("*"):

            if not file.is_file():
                continue

            suffix = file.suffix.lower()

            try:

                if suffix == ".txt":

                    text = file.read_text(
                        encoding="utf-8",
                        errors="ignore",
                    )

                elif suffix == ".docx":

                    document = Document(file)

                    text = "\n".join(
                        paragraph.text
                        for paragraph in document.paragraphs
                    )

                elif suffix == ".pdf":

                    reader = PdfReader(str(file))

                    pages = []

                    for page in reader.pages:

                        extracted = page.extract_text()

                        if extracted:
                            pages.append(extracted)

                    text = "\n".join(pages)

                else:
                    continue

            except Exception:
                continue

            if not text.strip():
                continue

            self.index.add(
                IndexedDocument(
                    title=file.name,
                    path=str(file),
                    text=text,
                )
            )

        return self.index
